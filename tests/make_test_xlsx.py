import openpyxl, datetime, sys
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Лист1"; ws["A1"]="x"
ws2=wb.create_sheet("Цены")
ws2.append(["Товар","Цена","Срок"])
ws2.append(["Молоко",42.5,datetime.date(2026,12,31)])
ws2.append(["Хлеб",31.9,datetime.date(2027,1,5)])
ws2.append(["Ряд 3 без цены",None,None])
wb.save(sys.argv[1])
